#!/usr/bin/env python3
# Script:   .claude/skills/ppt-creator/scripts/check_deck.py (research-to-policy)
# Inputs:   a .pptx
# Outputs:  a per-slide report on stdout; exit 1 if any ERROR
# Seed:     none
# Env:      python3, python-pptx
#
# Lints a deck against the slide grammar in ../SKILL.md:
#   title + ONE kind of content + source; <= 4 charts; no free text that reads
#   the chart; nothing under 11 pt. On gl_deck-built slides it also checks the
#   fixed format: Arial everywhere, title 24 pt, chart title 14 pt bold centred,
#   Notas / Fuente 12 pt, charts centred on the slide, no internal evidence
#   references, and — measured with Arial metrics — no text overflowing its box,
#   no blocks overlapping, nothing inside the footer band.
#
# Shapes written by gl_deck.py carry a role name ("gl:title", "gl:chart", ...).
# Shapes from any other builder are classified by placeholder type and shape
# kind, so the lint also runs on decks this skill did not build (a text box it
# cannot place is reported as free text — the thing the grammar forbids).
#
#   python3 check_deck.py deck.pptx            # report, exit 1 on errors
#   python3 check_deck.py deck.pptx --quiet    # only slides with findings

from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.util import Emu, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
import gl_deck as G  # noqa: E402  (format constants and Arial text measurement)

PT_MIN = 11
SIDE_MAX_WORDS = 45

# Layouts whose content is not a title/content/source slide.
FREE_LAYOUTS = {"Title Slide", "Title Slide With Background Image", "Statement Text",
                "Closing Slide", "Full Image"}

TITLE_ZONE_IN = 1.3     # a role-less text box starting above this is the headline
SUBTITLE_ZONE_IN = 2.0  # ...and the next one above this, the chart title

CONTENT_ROLES = {"gl:chart": "chart", "gl:table": "table", "gl:list": "list",
                 "gl:image": "image", "gl:statement": "text"}
TEXT_ROLES = {"gl:title", "gl:subtitle", "gl:source", "gl:note", "gl:side", "gl:meta",
              "gl:list-head", "gl:section-label", "gl:section-icon", "gl:section-sub",
              "gl:section-arrow", "gl:list-icon"}
BLOCK_ROLES = ("gl:title", "gl:subtitle", "gl:chart", "gl:image", "gl:table", "gl:note",
               "gl:source")
TOL_IN = 0.06
FOOTER_PH = {PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.DATE}


def role(shape) -> str:
    """Best guess at what a shape is doing on the slide."""
    if shape.name.startswith("gl:"):
        return shape.name
    if shape.is_placeholder:
        t = shape.placeholder_format.type
        if t in FOOTER_PH:
            return "footer"
        if t in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE):
            return "gl:title"
        if t == PP_PLACEHOLDER.SUBTITLE:
            return "gl:subtitle"
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        return "gl:chart"
    if getattr(shape, "has_table", False) and shape.has_table:
        return "gl:table"
    if getattr(shape, "has_chart", False) and shape.has_chart:
        return "gl:chart"
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        return "diagram"
    if shape.has_text_frame:
        txt = shape.text_frame.text.strip()
        if not txt:
            return "empty"
        if txt.lower().startswith(("fuente", "source")) or "Fuente:" in txt:
            return "gl:source"
        if len(txt) <= 5:          # page labels such as "4b", "12"
            return "footer"
        return "free-text"
    return "other"


def small_runs(shape):
    """Explicit run sizes under the floor (inherited sizes are the template's)."""
    frames = []
    if shape.has_text_frame:
        frames.append(shape.text_frame)
    if getattr(shape, "has_table", False) and shape.has_table:
        frames += [c.text_frame for row in shape.table.rows for c in row.cells]
    out = []
    for tf in frames:
        for p in tf.paragraphs:
            for r in p.runs:
                size = r.font.size or p.font.size
                if size is not None and size < Pt(PT_MIN) and r.text.strip():
                    out.append((size.pt, r.text.strip()[:40]))
    return out


def _box(sh):
    return tuple(Emu(v).inches for v in (sh.left, sh.top, sh.width, sh.height))


def _runs(sh):
    frames = []
    if sh.has_text_frame:
        frames.append(sh.text_frame)
    if getattr(sh, "has_table", False) and sh.has_table:
        frames += [c.text_frame for row in sh.table.rows for c in row.cells]
    for tf in frames:
        for p in tf.paragraphs:
            for r in p.runs:
                if r.text.strip():
                    yield p, r


def _sizes(sh):
    return {(r.font.size or p.font.size).pt for p, r in _runs(sh)
            if (r.font.size or p.font.size) is not None}


def check_format(slide, roles):
    """The fixed GL format, for slides whose shapes carry gl: roles. Title, statement
    and closing layouts keep the template's own sizes: only font and references apply."""
    errors = []
    free = slide.slide_layout.name.strip() in FREE_LAYOUTS
    blob = []
    for r, sh in roles:
        for p, run in _runs(sh):
            blob.append(run.text)
            if run.font.name not in (None, G.FONT):
                errors.append(f"font {run.font.name!r} (not {G.FONT}): “{run.text.strip()[:40]}”")
                break
    if slide.has_notes_slide:
        blob.append(slide.notes_slide.notes_text_frame.text)
    m = G.EVIDENCE_RE.search("\n".join(blob))
    if m:
        errors.append(f"internal evidence reference: “{m.group(0).strip()}”")
    want = {"gl:title": G.PT_TITLE, "gl:subtitle": G.PT_CHART_TITLE,
            "gl:source": G.PT_SOURCE, "gl:note": G.PT_NOTE}
    if free:
        return errors
    for r, sh in roles:
        if r in want and sh.has_text_frame and sh.text_frame.text.strip():
            bad = {x for x in _sizes(sh) if abs(x - want[r]) > 0.1}
            if bad or not _sizes(sh):
                errors.append(f"{r} at {sorted(bad) or 'inherited'} pt, not {want[r]} pt")
        if r == "gl:subtitle":
            if any(not run.font.bold for _, run in _runs(sh)):
                errors.append("chart title not bold")
            if any(p.alignment is not None and p.alignment != 2 for p in sh.text_frame.paragraphs):
                errors.append("chart title not centred")
        if r == "gl:side":
            errors.append("side text: the category definition goes in the Notas line "
                          "above the Fuente (note=...)")
        # text that does not fit its box, measured with Arial metrics
        if r in ("gl:title", "gl:subtitle", "gl:note", "gl:source") and sh.has_text_frame:
            txt = sh.text_frame.text
            if txt.strip():
                size = max(_sizes(sh) or {want.get(r, 12)})
                bold = any(run.font.bold for _, run in _runs(sh))
                l, t, w, h = _box(sh)
                need = G.block_h(txt, size, w, bold)
                if need > h + TOL_IN:
                    errors.append(f"{r} overflows its box ({need:.2f} in needed, {h:.2f} in): "
                                  f"“{txt.strip()[:50]}”")
                if r == "gl:title" and G.n_lines(txt, size, w, bold) > G.TITLE_MAX_LINES:
                    errors.append(f"title longer than {G.TITLE_MAX_LINES} lines")
                if r == "gl:subtitle" and G.n_lines(txt, size, w, bold) > G.CHART_TITLE_MAX_LINES:
                    errors.append(f"chart title longer than {G.CHART_TITLE_MAX_LINES} lines")
    # the chart (or grid of same-type charts) is centred horizontally on the slide
    vis = [_box(sh) for r, sh in roles if r in ("gl:chart", "gl:image")]
    if vis:
        left = min(b[0] for b in vis)
        right = max(b[0] + b[2] for b in vis)
        off = (left + right) / 2 - G.SLIDE_W / 2
        if abs(off) > TOL_IN:
            errors.append(f"chart off-centre horizontally by {off:+.2f} in")
    # blocks overlapping each other or entering the footer band
    blocks = [(r, _box(sh)) for r, sh in roles if r in BLOCK_ROLES and sh.width]
    for r, (l, t, w, h) in blocks:
        if t + h > G.FOOTER_TOP + 0.01:
            errors.append(f"{r} reaches into the footer band (bottom at {t + h:.2f} in)")
    for i in range(len(blocks)):
        for j in range(i + 1, len(blocks)):
            (ra, a), (rb, b) = blocks[i], blocks[j]
            if ra == rb and ra in ("gl:chart", "gl:image"):
                continue
            ox = min(a[0] + a[2], b[0] + b[2]) - max(a[0], b[0])
            oy = min(a[1] + a[3], b[1] + b[3]) - max(a[1], b[1])
            if ox > TOL_IN and oy > TOL_IN:
                errors.append(f"{ra} overlaps {rb} by {oy:.2f} in")
    return errors


def check_slide(slide):
    errors, warns = [], []
    layout = slide.slide_layout.name.strip()
    roles = [(role(sh), sh) for sh in slide.shapes]
    # Decks not built by gl_deck.py: promote the top text boxes to headline and
    # chart title by position, so only the text below them counts as free text.
    if not any(r == "gl:title" for r, _ in roles):
        loose = sorted((sh.top, k) for k, (r, sh) in enumerate(roles)
                       if r == "free-text" and sh.top is not None)
        for pos, (top, k) in enumerate(loose[:2]):
            zone = TITLE_ZONE_IN if pos == 0 else SUBTITLE_ZONE_IN
            if top < zone * 914400:
                roles[k] = ("gl:title" if pos == 0 else "gl:subtitle", roles[k][1])
    names = [r for r, _ in roles]

    for r, sh in roles:
        for pt, txt in small_runs(sh):
            errors.append(f"text at {pt:g} pt < {PT_MIN} pt: “{txt}”")

    if any(r.startswith("gl:") and sh.name.startswith("gl:") for r, sh in roles):
        errors += check_format(slide, roles)

    if layout in FREE_LAYOUTS:
        return errors, warns

    is_section = "gl:section-label" in names
    if "gl:title" not in names:
        errors.append("no title")
    if is_section:
        return errors, warns

    kinds = {CONTENT_ROLES[r] for r in names if r in CONTENT_ROLES}
    if "diagram" in names:
        kinds.add("diagram")
    if not kinds:
        errors.append("no content (chart, table, list, image or text)")
    elif len(kinds) > 1:
        errors.append(f"mixes content kinds: {sorted(kinds)} — one kind per slide")

    n_charts = names.count("gl:chart")
    if n_charts > 4:
        errors.append(f"{n_charts} charts — at most 4")
    elif n_charts > 1:
        warns.append(f"{n_charts} charts — ideally one per slide")

    if "gl:source" not in names:
        errors.append("no source line")
    if "chart" in kinds and "gl:subtitle" not in names:
        warns.append("chart without a chart title (the line naming the variables)")

    for r, sh in roles:
        if r == "free-text":
            errors.append(f"free text box (reads or explains the content?): "
                          f"“{sh.text_frame.text.strip()[:70]}”")
        if r == "gl:note" and len(sh.text_frame.text.split()) > SIDE_MAX_WORDS + 15:
            warns.append(f"Notas has {len(sh.text_frame.text.split())} words: define the "
                         "category, do not read the chart")
        if r == "gl:side":
            if "chart" not in kinds:
                errors.append("side text without a chart")
            n = len(sh.text_frame.text.split())
            if n > SIDE_MAX_WORDS:
                warns.append(f"side text has {n} words (> {SIDE_MAX_WORDS}): define the "
                             "category, do not read the chart")
    return errors, warns


def main(argv):
    if not argv or argv[0] in ("-h", "--help"):
        sys.exit(__doc__ or "usage: check_deck.py deck.pptx [--quiet]")
    path, quiet = Path(argv[0]), "--quiet" in argv
    prs = Presentation(str(path))
    n_err = n_warn = 0
    for i, slide in enumerate(prs.slides, 1):
        errors, warns = check_slide(slide)
        n_err += len(errors)
        n_warn += len(warns)
        if quiet and not (errors or warns):
            continue
        title = next((sh.text_frame.text.strip() for sh in slide.shapes
                      if role(sh) == "gl:title" and sh.has_text_frame), "")
        print(f"{i:3d} [{slide.slide_layout.name.strip()}] {title[:70]}")
        for e in errors:
            print(f"      ERROR  {e}")
        for w in warns:
            print(f"      warn   {w}")
    print(f"\n{len(prs.slides)} slides · {n_err} errors · {n_warn} warnings")
    sys.exit(1 if n_err else 0)


if __name__ == "__main__":
    main(sys.argv[1:])
